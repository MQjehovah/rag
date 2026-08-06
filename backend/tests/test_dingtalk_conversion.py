import io
import sys
from types import SimpleNamespace

import pytest
from openpyxl import Workbook

from app.config import settings
from app.core.dingtalk import DingTalkClient
from app.core.pdf_ocr import PDFOCRService


def test_pdf_extraction_uses_pypdf_fallback(monkeypatch):
    class FakePage:
        def extract_text(self):
            return "第一页内容"

    class FakeReader:
        def __init__(self, _stream):
            self.pages = [FakePage()]

    monkeypatch.setitem(
        sys.modules,
        "pypdf",
        SimpleNamespace(PdfReader=FakeReader),
    )

    assert DingTalkClient._extract_pdf_text(b"%PDF-test") == "第一页内容"


def test_xlsx_conversion_produces_valid_markdown_table():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "设备清单"
    sheet.append(["名称", "数量"])
    sheet.append(["机器人", 2])

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()

    markdown = DingTalkClient._extract_text(buffer.getvalue(), "xlsx")

    assert "## 设备清单" in markdown
    assert "| 名称 | 数量 |" in markdown
    assert "| --- | --- |" in markdown
    assert "| 机器人 | 2 |" in markdown


def test_pdf_markdown_keeps_pages_tables_and_visual_warning(monkeypatch):
    monkeypatch.setattr(settings, "pdf_ocr_enabled", False)
    pages = [
        {"page_number": 1, "text": "安装准备\nOperation\n准备工具", "image_count": 2},
        {"page_number": 2, "text": "1. 打开后盖\nPower On - Connect Battery\nOperation\n连接电源线", "image_count": 1},
        {"page_number": 3, "text": "启动设备\nOperation\n打开开关", "image_count": 0},
    ]
    tables = {
        2: ["| 接口 | 作用 |\n| --- | --- |\n| 电源口 | 供电 |"],
    }
    monkeypatch.setattr(
        DingTalkClient,
        "_extract_pdf_pages",
        staticmethod(lambda _content: pages),
    )
    monkeypatch.setattr(
        DingTalkClient,
        "_extract_pdf_tables",
        staticmethod(lambda _content: tables),
    )

    markdown = DingTalkClient._pdf_to_markdown(b"%PDF-test")

    assert "共 3 页，检测到 3 个图片对象" in markdown
    assert "## 第 1 页 - 安装准备" in markdown
    assert "## 第 2 页 - Power On - Connect Battery" in markdown
    assert "1. 打开后盖" in markdown
    assert "### 本页识别表格" in markdown
    assert "| 接口 | 作用 |" in markdown
    assert "Operation" not in markdown


def test_pdf_markdown_appends_only_novel_ocr_text(monkeypatch):
    pages = [
        {
            "page_number": 1,
            "text": "Connect Battery\nTurn on the battery power",
            "image_count": 1,
        },
    ]
    monkeypatch.setattr(settings, "pdf_ocr_enabled", True)
    monkeypatch.setattr(
        DingTalkClient,
        "_extract_pdf_pages",
        staticmethod(lambda _content: pages),
    )
    monkeypatch.setattr(
        DingTalkClient,
        "_extract_pdf_tables",
        staticmethod(lambda _content: {}),
    )
    monkeypatch.setattr(
        PDFOCRService,
        "extract_novel_text",
        classmethod(lambda _cls, _content, _pages: ({1: ["Cleaner Login"]}, 1)),
    )

    markdown = DingTalkClient._pdf_to_markdown(b"%PDF-test")

    assert "已对 1 页执行图片文字识别，补充 1 行" in markdown
    assert "### 本页图片文字（OCR）" in markdown
    assert "- Cleaner Login" in markdown
    assert "可能存在少量误差" in markdown


def test_pdf_ocr_novel_line_deduplicates_text_layer(monkeypatch):
    monkeypatch.setattr(settings, "pdf_ocr_min_novel_chars", 4)
    monkeypatch.setattr(settings, "pdf_ocr_duplicate_similarity", 0.86)
    source = (
        "Install the battery and connect the wiring\n"
        "Turn on the battery power"
    )

    assert not PDFOCRService._is_novel_line(
        "Install the battery and connect the wiring;",
        source,
    )
    assert not PDFOCRService._is_novel_line("123", source)
    assert PDFOCRService._is_novel_line("Cleaner Login", source)


def test_pdf_replacement_character_is_detected():
    assert PDFOCRService.has_replacement_chars("修复时长：��小时")
    assert not PDFOCRService.has_replacement_chars("修复时长：48小时")


def test_pdf_ocr_corrects_terms_merges_fragments_and_removes_noise():
    cleaned = PDFOCRService._merge_and_correct_lines([
        "Basemen",
        "Mimsion report",
        "deceler",
        "ation zo",
        "unidirecti",
        "onal",
        "ROSIWVIT",
        "Connected, secure",
        "Cleaner Login",
    ])

    assert "Basement A" in cleaned
    assert "Mission report" in cleaned
    assert "Deceleration zone" in cleaned
    assert "Unidirectional" in cleaned
    assert "ROSIWVIT" not in cleaned
    assert "Connected, secure" not in cleaned
    assert "Cleaner Login" in cleaned


def test_pdf_table_text_is_removed_when_markdown_table_exists():
    body = [
        "Name",
        "Application Scenarios",
        "Virtual Wall",
        "Narrative text outside the table",
    ]
    tables = [
        "| Name | Application Scenarios |\n"
        "| --- | --- |\n"
        "| Virtual Wall | Robots cannot enter |"
    ]

    cleaned = DingTalkClient._deduplicate_table_lines(body, tables)

    assert cleaned == ["Narrative text outside the table"]


def test_short_percentages_are_removed_when_table_already_contains_them():
    body = ["延保说明", "15%", "23%", "普通正文"]
    tables = [
        "| 产品 | 第1年 | 第2年 |\n"
        "| --- | --- | --- |\n"
        "| SC 50 | 15% | 23% |"
    ]

    cleaned = DingTalkClient._deduplicate_table_lines(body, tables)

    assert cleaned == ["延保说明", "普通正文"]


def test_markitdown_table_block_is_removed_when_enhanced_table_exists():
    body = [
        "| Name | Application Scenarios |",
        "| --- | --- |",
        "| Virtual Wall | Robots cannot enter |",
        "Narrative text outside the table",
    ]
    tables = [
        "| Name | Application Scenarios |\n"
        "| --- | --- |\n"
        "| Virtual Wall | Robots cannot enter |"
    ]

    cleaned = DingTalkClient._deduplicate_table_lines(body, tables)

    assert cleaned == ["Narrative text outside the table"]


def test_broken_markitdown_table_and_fragments_are_removed():
    body = [
        "Intro outside the table",
        "| Tro | u b | l e sh | o o | t i n g | | |",
        "| | Poor clean | | | | Clean hose | |",
        "| --- | --- | --- | --- | --- | --- | --- |",
        "No clean water Please check tank assembled Put tank properly",
        "4",
        "| | pump out | | properly | | assembled | |",
        "| --- | --- | --- | --- | --- | --- | --- |",
        "Conclusion outside the table",
    ]
    tables = [
        "| Failure | Possible reason | Solution |\n"
        "| --- | --- | --- |\n"
        "| No clean water pump out | Check tank | Put tank properly |"
    ]

    cleaned = DingTalkClient._deduplicate_table_lines(body, tables)

    assert cleaned == ["Intro outside the table", "Conclusion outside the table"]


def test_detached_table_text_is_removed_by_word_overlap():
    body = [
        "Narrative text outside the table",
        "No clean water Please check tank assembled Put tank properly",
        "4",
    ]
    tables = [
        "| No. | Failure | Possible reason | Solution |\n"
        "| --- | --- | --- | --- |\n"
        "| 4 | No clean water | Please check tank assembled | Put tank properly |"
    ]

    cleaned = DingTalkClient._deduplicate_table_lines(body, tables)

    assert cleaned == ["Narrative text outside the table"]


def test_fragmented_table_heading_is_recovered():
    title = DingTalkClient._recover_fragmented_table_heading([
        "squeegee；",
        "| Tro | u b | l e sh | o o | t i n g | | |",
    ])

    assert title == "Troubleshooting"


def test_corrupted_table_is_rebuilt_from_ocr_coordinates():
    class FakeTable:
        bbox = (0, 0, 100, 40)
        cells = [
            (0, 0, 50, 20),
            (50, 0, 100, 20),
            (0, 20, 50, 40),
            (50, 20, 100, 40),
        ]

    rebuilt = DingTalkClient._rebuild_table_from_ocr(
        [["城市", "��小时"], ["北京", "��小时"]],
        FakeTable(),
        [
            {"text": "城市", "bbox": (5, 5, 20, 15)},
            {"text": "修复时长", "bbox": (55, 5, 90, 15)},
            {"text": "北京", "bbox": (5, 25, 20, 35)},
            {"text": "48小时", "bbox": (55, 25, 85, 35)},
        ],
    )

    assert rebuilt == [["城市", "修复时长"], ["北京", "48小时"]]


def test_vertical_merged_cell_text_is_assigned_to_first_row():
    class FakeTable:
        bbox = (0, 0, 100, 80)
        cells = [
            (0, 0, 50, 20),
            (50, 0, 100, 20),
            (0, 20, 50, 80),
            (50, 20, 100, 40),
            (50, 40, 100, 60),
            (50, 60, 100, 80),
        ]

    rebuilt = DingTalkClient._rebuild_table_from_ocr(
        [["产品", "等级"], ["", "大修"], ["SC50", "中修"], ["", "小修"]],
        FakeTable(),
        [
            {"text": "产品", "bbox": (5, 5, 20, 15)},
            {"text": "等级", "bbox": (55, 5, 75, 15)},
            {"text": "SC50", "bbox": (10, 43, 35, 55)},
            {"text": "大修", "bbox": (55, 23, 75, 35)},
            {"text": "中修", "bbox": (55, 43, 75, 55)},
            {"text": "小修", "bbox": (55, 63, 75, 75)},
        ],
    )

    assert rebuilt[1][0] == "SC50"
    assert rebuilt[2][0] == ""


def test_pdf_table_merge_cells_preserve_rowspan_and_colspan():
    class FakeTable:
        cells = [
            (0, 0, 40, 20),
            (40, 0, 100, 20),
            (0, 20, 40, 60),
            (40, 20, 100, 40),
            (40, 40, 100, 60),
            (0, 60, 100, 80),
        ]

    merges = DingTalkClient._table_merge_cells(
        FakeTable(),
        row_count=4,
        column_count=2,
        active_columns=[0, 1],
    )

    assert {"row": 1, "col": 0, "rowspan": 2, "colspan": 1} in merges
    assert {"row": 3, "col": 0, "rowspan": 1, "colspan": 2} in merges


def test_corrupted_pdf_page_uses_full_page_ocr(monkeypatch):
    pages = [{
        "page_number": 1,
        "text": "服务响应\n修复时长：��小时",
        "image_count": 0,
    }]
    monkeypatch.setattr(
        DingTalkClient,
        "_extract_pdf_pages",
        staticmethod(lambda _content: pages),
    )
    monkeypatch.setattr(
        PDFOCRService,
        "extract_corrupted_pages",
        classmethod(lambda _cls, _content, _pages: (
            {1: ["服务响应", "修复时长：48小时"]},
            {1: [{"text": "48小时", "bbox": (50, 20, 90, 30)}]},
            1,
        )),
    )
    monkeypatch.setattr(
        DingTalkClient,
        "_extract_pdf_tables",
        staticmethod(lambda _content, _detections: {}),
    )

    markdown = DingTalkClient._pdf_to_markdown(b"%PDF-test")

    assert "修复时长：48小时" in markdown
    assert "�" not in markdown
    assert "检测到 1 页字体编码异常，已通过整页OCR恢复 1 页" in markdown


def test_table_only_pdf_does_not_show_empty_body_placeholder(monkeypatch):
    pages = [{
        "page_number": 1,
        "text": "squeegee；\n| Tro | u b | l e sh | o o | t i n g |",
        "image_count": 0,
    }]
    tables = {
        1: [
            "| Failure | Possible reason | Solution |\n"
            "| --- | --- | --- |\n"
            "| Poor clean performance | squeegee dirty | Clean squeegee |"
        ],
    }
    monkeypatch.setattr(
        DingTalkClient,
        "_extract_pdf_pages",
        staticmethod(lambda _content: pages),
    )
    monkeypatch.setattr(
        DingTalkClient,
        "_extract_pdf_tables",
        staticmethod(lambda _content: tables),
    )

    markdown = DingTalkClient._pdf_to_markdown(b"%PDF-test")

    assert "## 第 1 页 - Troubleshooting" in markdown
    assert "### 本页识别表格" in markdown
    assert "本页未提取到除标题外的文字" not in markdown


def test_markitdown_heading_marker_is_removed_from_pdf_page_title():
    title, index = DingTalkClient._pdf_page_title(
        ["## Robot Deployment", "Body"],
        1,
    )

    assert title == "Robot Deployment"
    assert index == 0


def test_numbered_section_heading_is_preferred_over_wrapped_body_text():
    title, index = DingTalkClient._pdf_page_title(
        [
            "由公司销售或授权经销商销售、租赁的设备（以下",
            "1.适用产品",
            "简称产品。",
        ],
        2,
    )

    assert title == "1.适用产品"
    assert index == 1


def test_pdf_table_is_not_mistaken_for_page_title():
    title, index = DingTalkClient._pdf_page_title(
        ["| Name | Value |", "| --- | --- |", "| A | B |"],
        2,
    )

    assert title == "第 2 页"
    assert index == -1
